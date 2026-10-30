"""
นำเข้าบัญชีแอดมินหลายคนจากไฟล์ Excel
SuperAdmin ดาวน์โหลด template → กรอก 1 แถวต่อ 1 คน → อัปโหลดกลับ ระบบสร้างรหัสผ่านเริ่มต้นให้ทุกคน
แล้วคืนรายชื่อพร้อมรหัสผ่านกลับไปครั้งเดียว (ไม่เก็บรหัสจริงไว้ที่ไหน) ให้ SuperAdmin แจ้งแต่ละคน

- ตรวจทุกแถวก่อน — มีแถวผิดแม้แถวเดียวจะไม่สร้างบัญชีใดเลย (แก้ไฟล์แล้วอัปโหลดใหม่ทั้งไฟล์ได้ ไม่ต้องไล่หาว่าคนไหนสร้างไปแล้ว)
- บัญชีที่นำเข้ามี password_changed_at = NULL → ล็อกอินครั้งแรกถูกบังคับให้เปลี่ยนรหัสผ่านเอง
  (รหัสเริ่มต้นผ่านมือ SuperAdmin มาแล้ว ไม่ควรใช้ต่อ)
"""
import io
import secrets
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.services.department_scope import SUPER_ADMIN_ROLE
from app.services.password_policy import password_problems

MAX_ROWS = 200
ROLES = ("admin", SUPER_ADMIN_ROLE)

# หัวคอลัมน์ใน template — อ่านกลับด้วยคำในวงเล็บ จึงแก้คำไทยหน้าวงเล็บได้โดยไม่กระทบการนำเข้า
COLUMNS = [
    ("ชื่อแสดงผล (display_name)", "display_name", 28),
    ("Username (username)", "username", 22),
    ("หมวดงาน / แผนก (department)", "department", 30),
    ("บทบาท (role)", "role", 26),
]

# ตัดตัวที่อ่านสับสนเวลาบอกรหัสกัน (0/O, 1/l/I) ออก — สัญลักษณ์เลือกเฉพาะที่พิมพ์ง่ายบนแป้นทั่วไป
_UPPER = "ABCDEFGHJKLMNPQRSTUVWXYZ"
_LOWER = "abcdefghijkmnopqrstuvwxyz"
_DIGITS = "23456789"
_SYMBOLS = "!@#$%&*?"
PASSWORD_LENGTH = 12


def generate_password(username: Optional[str] = None) -> str:
    """รหัสผ่านสุ่ม 12 ตัว ผ่านกฎ password_policy ครบทุกข้อ"""
    pool = _UPPER + _LOWER + _DIGITS + _SYMBOLS
    while True:
        chars = [secrets.choice(s) for s in (_UPPER, _LOWER, _DIGITS, _SYMBOLS)]
        chars += [secrets.choice(pool) for _ in range(PASSWORD_LENGTH - len(chars))]
        secrets.SystemRandom().shuffle(chars)
        password = "".join(chars)
        if not password_problems(password, username):
            return password


def build_template(departments: List[str]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "รายชื่อแอดมิน"
    header_fill = PatternFill("solid", fgColor="4B2E83")
    for col, (title, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    # แถวตัวอย่าง — ตอนนำเข้าข้ามแถวที่ username ขึ้นต้นด้วย "ตัวอย่าง" ไม่ต้องลบก่อนอัปโหลด
    example = ["สมชาย ใจดี", "ตัวอย่าง_somchai", departments[0] if departments else "งานสารสนเทศ", "admin"]
    for col, value in enumerate(example, start=1):
        ws.cell(row=2, column=col, value=value).font = Font(italic=True, color="888888")

    role_rule = DataValidation(type="list", formula1='"' + ",".join(ROLES) + '"', allow_blank=True)
    role_rule.error = "เลือก admin หรือ System Administrator"
    ws.add_data_validation(role_rule)
    role_rule.add(f"D2:D{MAX_ROWS + 1}")

    guide = wb.create_sheet("วิธีกรอก")
    guide.column_dimensions["A"].width = 100
    lines = [
        "วิธีกรอกไฟล์นำเข้าบัญชีแอดมิน",
        "",
        "1. กรอก 1 แถวต่อ 1 คน ในชีต \"รายชื่อแอดมิน\" (สูงสุด %d คนต่อไฟล์) — แถวตัวอย่างสีเทาไม่ต้องลบ ระบบข้ามให้" % MAX_ROWS,
        "2. ชื่อแสดงผล และ Username ต้องกรอกทุกแถว — Username ใช้ a-z, 0-9, จุด, ขีด, ขีดล่าง และห้ามซ้ำกับที่มีอยู่",
        "3. หมวดงาน/แผนก ต้องกรอกสำหรับบทบาท admin — สะกดให้ตรงกับแผนกเดิมในระบบ (รายชื่อแผนกด้านล่าง)",
        "4. บทบาท: admin = แอดมินแผนก (เห็นเฉพาะข้อมูลแผนกตัวเอง), System Administrator = SuperAdmin — เว้นว่าง = admin",
        "5. ไม่ต้องกรอกรหัสผ่าน — ระบบสร้างรหัสผ่านเริ่มต้นให้ และแสดงผลให้ดาวน์โหลดหลังนำเข้า (แสดงครั้งเดียว)",
        "6. ผู้ใช้ต้องเปลี่ยนรหัสผ่านเองตอนเข้าสู่ระบบครั้งแรก",
        "7. ถ้ามีแถวใดผิด ระบบจะไม่สร้างบัญชีใดเลย — แก้ตามที่แจ้งแล้วอัปโหลดไฟล์เดิมใหม่ได้ทันที",
        "",
        "แผนกที่มีอยู่ในระบบตอนนี้:",
    ] + ["   • " + d for d in departments]
    for i, line in enumerate(lines, start=1):
        guide.cell(row=i, column=1, value=line)
    guide["A1"].font = Font(bold=True, size=14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_rows(content: bytes) -> Tuple[List[Dict], List[Dict]]:
    """อ่านไฟล์ → (rows, errors) — rows มี row (เลขแถวใน Excel), display_name, username, department, role
    ตรวจเฉพาะเนื้อหาในไฟล์ (รวมซ้ำกันเองในไฟล์) — การซ้ำกับบัญชีใน DB ตรวจที่ router"""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise ValueError("อ่านไฟล์ไม่ได้ — กรุณาใช้ไฟล์ .xlsx จาก template ของระบบ")
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header = [_clean(h) for h in next(rows_iter, [])]

    index = {}
    for key in (c[1] for c in COLUMNS):
        for i, h in enumerate(header):
            if f"({key})" in h or h == key:
                index[key] = i
    missing = [c[0] for c in COLUMNS if c[1] not in index and c[1] != "role"]
    if missing:
        raise ValueError("ไม่พบคอลัมน์: " + ", ".join(missing) + " — กรุณาใช้ template ของระบบ")

    rows, errors, seen = [], [], {}
    for excel_row, values in enumerate(rows_iter, start=2):
        get = lambda k: _clean(values[index[k]]) if k in index and index[k] < len(values) else ""
        item = {k: get(k) for k in ("display_name", "username", "department", "role")}
        if not any(item.values()) or item["username"].startswith("ตัวอย่าง"):
            continue
        if len(rows) + len(errors) >= MAX_ROWS:
            raise ValueError(f"นำเข้าได้ครั้งละไม่เกิน {MAX_ROWS} คน")

        problems = []
        username = item["username"]
        if not item["display_name"]:
            problems.append("ไม่ได้กรอกชื่อแสดงผล")
        if not username:
            problems.append("ไม่ได้กรอก Username")
        elif not all(ch.isascii() and (ch.isalnum() or ch in "._-") for ch in username) or len(username) > 100:
            problems.append("Username ใช้ได้เฉพาะ a-z, 0-9, จุด, ขีด, ขีดล่าง")
        elif username.lower() in seen:
            problems.append(f"Username ซ้ำกับแถว {seen[username.lower()]} ในไฟล์")
        if username:
            seen.setdefault(username.lower(), excel_row)

        role = item["role"] or "admin"
        match = next((r for r in ROLES if r.lower() == role.lower()), None)
        if not match:
            problems.append("บทบาทต้องเป็น admin หรือ System Administrator")
        item["role"] = match or role
        if item["role"] != SUPER_ADMIN_ROLE and not item["department"]:
            problems.append("แอดมินแผนกต้องระบุหมวดงาน/แผนก")
        if len(item["display_name"]) > 100 or len(item["department"]) > 100:
            problems.append("ชื่อแสดงผล/แผนกยาวเกิน 100 ตัวอักษร")

        if problems:
            errors.append({"row": excel_row, "username": username, "message": ", ".join(problems)})
        else:
            rows.append({"row": excel_row, **item, "department": item["department"] or None})
    wb.close()
    return rows, errors
