"""
Tests: นำเข้าบัญชีแอดมินหลายคนจาก Excel (POST /api/auth/users/import, GET .../import/template)
- เฉพาะ SuperAdmin, สร้างรหัสผ่านเริ่มต้นที่ผ่านกฎ password_policy และคืนมาให้แจ้งผู้ใช้
- มีแถวผิดแม้แถวเดียว → ไม่สร้างบัญชีใดเลย
- บัญชีที่นำเข้าต้องเปลี่ยนรหัสผ่านตอนล็อกอินครั้งแรก
"""
import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.models.models import User
from app.services.password_policy import password_problems
from app.services.user_import import build_template, generate_password
from tests.conftest import auth_header

pytestmark = pytest.mark.asyncio

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sheet(rows):
    """template ของระบบ (มีแถวตัวอย่าง) + rows ที่ต่อท้าย"""
    wb = load_workbook(io.BytesIO(build_template(["งานสารสนเทศ"])))
    ws = wb.worksheets[0]
    for r in rows:
        ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _upload(client, username, content):
    return await client.post("/api/auth/users/import", headers=auth_header(username),
                             files={"file": ("admins.xlsx", content, XLSX)})


async def test_generated_passwords_pass_policy():
    for _ in range(200):
        assert password_problems(generate_password("somchai"), "somchai") == []


async def test_import_creates_users_and_returns_initial_passwords(client, admin_user, staff_user, db_session):
    # staff_user อยู่ งานสารสนเทศ → แผนกนี้มีอยู่แล้ว ไม่นับเป็นแผนกใหม่
    resp = await _upload(client, admin_user.username, _sheet([
        ("สมชาย ใจดี", "somchai", "งานสารสนเทศ", "admin"),
        ("สมหญิง รักงาน", "somying", "งานการพยาบาล", ""),          # role ว่าง = admin
        ("หัวหน้าระบบ", "boss", "", "System Administrator"),       # SuperAdmin ไม่ต้องมีแผนก
    ]))
    assert resp.status_code == 200
    body = resp.json()
    assert body["errors"] == []
    assert [u["username"] for u in body["created"]] == ["somchai", "somying", "boss"]
    assert body["new_departments"] == ["งานการพยาบาล"]

    # ล็อกอินด้วยรหัสเริ่มต้นได้ แต่ถูกบังคับเปลี่ยนรหัสก่อนใช้งาน
    pw = body["created"][0]["password"]
    login = (await client.post("/api/auth/login", json={"username": "somchai", "password": pw})).json()
    assert login["password_expired"] is True
    assert login["role"] == "admin" and login["department"] == "งานสารสนเทศ"
    assert (await db_session.execute(select(User).where(User.username == "somying"))).scalar_one().role == "admin"


async def test_any_invalid_row_creates_nobody(client, admin_user, staff_user, db_session):
    resp = await _upload(client, admin_user.username, _sheet([
        ("ok", "newone", "งานสารสนเทศ", "admin"),
        ("ไม่มีแผนก", "nodept", "", "admin"),
        ("ซ้ำในระบบ", staff_user.username, "งานสารสนเทศ", "admin"),
        ("ซ้ำในไฟล์", "NEWONE", "งานสารสนเทศ", "admin"),
        ("ผิดบทบาท", "badrole", "งานสารสนเทศ", "superuser"),
        ("", "ชื่อไทย", "งานสารสนเทศ", "admin"),
    ]))
    assert resp.status_code == 200
    body = resp.json()
    assert body["created"] == []
    assert [e["row"] for e in body["errors"]] == [4, 5, 6, 7, 8]
    assert (await db_session.execute(select(User).where(User.username == "newone"))).scalar_one_or_none() is None


async def test_import_requires_super_admin(client, staff_user):
    resp = await _upload(client, staff_user.username, _sheet([("a", "aaa", "งานสารสนเทศ", "admin")]))
    assert resp.status_code == 403
    resp = await client.get("/api/auth/users/import/template", headers=auth_header(staff_user.username))
    assert resp.status_code == 403


async def test_template_download_and_bad_file(client, admin_user):
    resp = await client.get("/api/auth/users/import/template", headers=auth_header(admin_user.username))
    assert resp.status_code == 200
    assert load_workbook(io.BytesIO(resp.content)).sheetnames == ["รายชื่อแอดมิน", "วิธีกรอก"]

    resp = await _upload(client, admin_user.username, b"not an excel file")
    assert resp.status_code == 400
