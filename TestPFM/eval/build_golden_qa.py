"""
แปลง Excel golden Q&A (ISO policy) ที่ทีมเตรียมไว้ ให้เป็น golden_qa.csv/.json
สำหรับใช้ประเมิน RAG answer quality (pillar 2 ใน TestPFM)

รูปแบบไฟล์ต้นทาง (sheet เดียวชื่อ "คำถาม"):
    แถว header ต่อ 1 หมวด: [category_name, None, "คำตอบ"/"คำตอบที่คาดว่าจะได้รับ", ...]
    แถวคำถามถัดมา 5 แถว:   [ลำดับ(float), question, expected_answer, ...]
    วนซ้ำแบบนี้ทีละหมวด

รัน:
    python TestPFM/eval/build_golden_qa.py "C:\\Users\\ITS\\Downloads\\Test_Chatbot iso.xlsx"

Output:
    TestPFM/eval/golden_qa.csv
    TestPFM/eval/golden_qa.json
"""
import csv
import json
import sys
from pathlib import Path

import openpyxl

OUT_DIR = Path(__file__).resolve().parent


def parse(xlsx_path: str):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = []
    current_category = None
    qid = 0

    for row in ws.iter_rows(min_row=1, values_only=True):
        col_a, col_b, col_c = row[0], row[1], row[2]

        if col_a is None and col_b is None and col_c is None:
            continue  # แถวว่าง

        # แถว header ของหมวด: col_a เป็นข้อความ (ไม่ใช่ตัวเลขลำดับ), col_b ว่าง
        if isinstance(col_a, str) and col_b is None:
            current_category = col_a.strip().replace("\n", " ")
            continue

        # แถวคำถาม: col_a เป็นเลขลำดับ, col_b = question, col_c = expected_answer
        if isinstance(col_a, (int, float)) and col_b:
            qid += 1
            rows.append({
                "id": f"iso-{qid:03d}",
                "category": current_category or "",
                "question": str(col_b).strip(),
                "expected_answer": str(col_c or "").strip(),
            })

    return rows


def main():
    if len(sys.argv) != 2:
        print("usage: python build_golden_qa.py <path-to-xlsx>")
        sys.exit(1)

    rows = parse(sys.argv[1])

    csv_path = OUT_DIR / "golden_qa.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=["id", "category", "question", "expected_answer"])
        writer.writeheader()
        writer.writerows(rows)

    json_path = OUT_DIR / "golden_qa.json"
    json_path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")

    categories = sorted(set(r["category"] for r in rows))
    print(f"แปลงสำเร็จ: {len(rows)} คำถาม จาก {len(categories)} หมวด")
    for c in categories:
        n = sum(1 for r in rows if r["category"] == c)
        print(f"  - {c} ({n})")
    print(f"\nบันทึกที่:\n  {csv_path}\n  {json_path}")


if __name__ == "__main__":
    main()
