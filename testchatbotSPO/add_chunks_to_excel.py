#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
TUH Chatbot AI — Enrich Excel with Admin History Chunks & Question Categories
=============================================================================
ดึงประวัติการตอบจากตาราง history ในฐานข้อมูลหลังบ้าน (Admin Chatbot History)
เพื่อดึง Chunk IDs, เอกสารอ้างอิงของ Chunk, และหมวดหมู่หัวข้อคำถาม
มารวมเข้ากับตารางผลการทดสอบ Excel
"""

import os
import sys
import json
import ssl
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pymysql
from dotenv import load_dotenv

# บังคับใช้ UTF-8 stdout บน Windows Console
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent

load_dotenv(PROJECT_ROOT / "Backend" / ".env")
load_dotenv(PROJECT_ROOT / ".env")


def main():
    print("=" * 75)
    print("🔍 กำลังดึงข้อมูล Chunk และหัวข้อคำถามจาก Admin History...")
    print("=" * 75)

    # 1. โหลด Category จาก questions_dataset.csv
    dataset_path = CURRENT_DIR / "questions_dataset.csv"
    if not dataset_path.exists():
        dataset_path = PROJECT_ROOT / "TestPFM" / "eval" / "golden_qa.csv"

    df_q = pd.read_csv(dataset_path)
    cat_map = dict(zip(df_q["question"].str.strip(), df_q["category"].str.strip()))
    print(f"✅ โหลดหัวข้อคำถาม (Category) สำเร็จ: {len(cat_map)} รายการ")

    # 2. โหลด Chunk map จาก sample_chunks.json
    chunks_file = PROJECT_ROOT / "sample_chunks.json"
    chunk_info_map = {}
    if chunks_file.exists():
        with open(chunks_file, "r", encoding="utf-8") as f:
            c_data = json.load(f)
            for item in c_data:
                cid = item.get("chunk_id")
                meta = item.get("metadata", {})
                chunk_info_map[cid] = {
                    "source": meta.get("source", ""),
                    "page": meta.get("page", 1)
                }
        print(f"✅ โหลดฐานข้อมูล Chunk สำเร็จ: {len(chunk_info_map)} chunks")

    # 3. เชื่อมต่อฐานข้อมูลเพื่อดึงประวัติ ChatHistory
    print("🔌 กำลังเชื่อมต่อฐานข้อมูล TiDB เพื่อดึงประวัติการแชท (history)...")
    ctx = ssl.create_default_context()
    conn = pymysql.connect(
        host=os.getenv("DB_HOST"),
        port=int(os.getenv("DB_PORT", 4000)),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD"),
        database=os.getenv("DB_NAME"),
        ssl=ctx,
        cursorclass=pymysql.cursors.DictCursor
    )

    with conn.cursor() as cur:
        cur.execute(
            "SELECT id, query, chunk_ids, referenced_docs, timestamp "
            "FROM history ORDER BY timestamp DESC LIMIT 500"
        )
        history_rows = cur.fetchall()
    conn.close()
    print(f"✅ โหลดประวัติการสนทนาจากฐานข้อมูลสำเร็จ: {len(history_rows)} รายการ")

    # จัดกลุ่มประวัติล่าสุดตามข้อความคำถาม
    history_map = {}
    for h in history_rows:
        q = (h["query"] or "").strip()
        if q and q not in history_map:
            history_map[q] = h

    # 4. โหลดผลการทดสอบ Excel เดิม
    source_excel = CURRENT_DIR / "test_result_gemma_all_80_questions.xlsx"
    if not source_excel.exists():
        # ค้นหาไฟล์ผลลัพธ์ล่าสุด
        xls_files = sorted(CURRENT_DIR.glob("test_result_gemma_*.xlsx"))
        if xls_files:
            source_excel = xls_files[-1]

    wb_orig = openpyxl.load_workbook(source_excel)
    ws_orig = wb_orig.active
    print(f"📖 กำลังอ่านผลการทดสอบจาก: {source_excel.name} ({ws_orig.max_row - 1} ข้อ)")

    # แผนที่แปลงชื่อไฟล์ PDF เป็น "หัวข้อสีชมพูในชีท" (ชื่อทางการของนโยบาย ISO)
    pdf_to_pink_title = {
        "POL-01-191-01_ISMS Framework_V.0.pdf": "กรอบการดําเนินงานระบบ ISMS (ISMS Framework)",
        "POL-01-191-02_IS Policy_V.0.pdf": "นโยบายบริหารความมั่นคงปลอดภัยสารสนเทศ (Information Security Policy)",
        "POL-01-191-03_Mobile device_V.0.pdf": "นโยบายคอมพิวเตอร์แบบพกพา (Mobile device)",
        "POL-01-191-04_Asset management policy_V.0.pdf": "นโยบายบริหารจัดการทรัพย์สินสารสนเทศ (Asset Management Policy)",
        "POL-01-191-05_Access control Policy_V.0.pdf": "นโยบายควบคุมการเข้าถึง (Access Control Policy)",
        "POL-01-191-06_Cryptographic Policy_V.0.pdf": "นโยบายการเข้ารหัสข้อมูล (Cryptographic Policy)",
        "POL-01-191-07_Clear desk and clear screen policy_V.0.pdf": "นโยบายโต๊ะทํางานปลอดเอกสารสําคัญ และนโยบายการป้องกันหน้าจอคอมพิวเตอร์  (Clear desk and clear screen policy)",
        "POL-01-191-08_Backup policy_V.0.pdf": "นโยบายการสํารองข้อมูล (Backup Policy)",
        "POL-01-191-09_Information transfer Policy_V.0.pdf": "นโยบายการถ่ายโอนสารสนเทศ (Information transfer policy)",
        "POL-01-191-10_Secure Development Policy_V.0.pdf": "นโยบายการพัฒนาระบบอย่างมั่นคงปลอดภัย (Secure Development Policy)",
        "POL-01-191-11_Supplier relationship policy_V.0.pdf": "นโยบายความสัมพันธ์กับผู้ให้บริการภายนอก (Supplier relationships policy)",
        "POL-01-191-12_Compliance Policy_V.0.pdf": "นโยบายการปฏิบัติตามกฎหมาย ระเบียบ ข้อบังคับ (Compliance Policy)",
        "POL-01-191-13_Teleworking policy_V.0.pdf": "นโยบายการปฏิบัติงานจากระยะไกล (Teleworking policy)",
        "POL-01-191-14_Baseline Configuration Policy_V.0.pdf": "นโยบายการตั้งค่าระบบตามมาตรฐาน (Baseline Configuration Policy)",
        "POL-01-191-15_Business Continuity Management Policy_V.0.pdf": "นโยบายการบริหารจัดการความต่อเนื่องในการดําเนินงานขององค์กร  (Procedure for Business Continuity Management)",
        "POL-01-191-16_Physical and Environmental policy_V.0.pdf": "นโยบายความปลอดภัยทางกายภาพ (Physical and environmental policy)",
    }

    # 5. สร้าง Workbook ใหม่ที่มีคอลัมน์ครบถ้วน
    wb_new = openpyxl.Workbook()
    ws_new = wb_new.active
    ws_new.title = "ผลการทดสอบแชทบอท"

    headers = [
        "ข้อที่",
        "หัวข้อของคำถาม (ในชีท)",
        "คำถาม",
        "คำตอบที่บอทตอบ",
        "คำตอบที่คาดหวัง",
        "Chunk ที่ใช้ตอบ (Chunk IDs)",
        "หัวข้อเอกสารและหน้าที่ Chunk สังกัด (หัวข้อสีชมพู)",
        "เวลาในการตอบ"
    ]
    ws_new.append(headers)

    for row_idx in range(2, ws_orig.max_row + 1):
        item_no = ws_orig.cell(row_idx, 1).value
        q = str(ws_orig.cell(row_idx, 2).value or "").strip()
        ans = ws_orig.cell(row_idx, 3).value
        exp = ws_orig.cell(row_idx, 4).value
        resp_time = ws_orig.cell(row_idx, 5).value

        # หัวข้อของคำถาม (Category)
        category = cat_map.get(q, "-")

        # ข้อมูล Chunk จากประวัติหลังบ้าน
        h = history_map.get(q, {})
        c_ids_raw = h.get("chunk_ids", "") or ""

        c_ids_list = [int(x) for x in c_ids_raw.split(",") if x.strip().isdigit()]
        c_ids_display = ", ".join([f"#{cid}" for cid in c_ids_list]) if c_ids_list else "-"

        # เอกสารและหน้าที่ Chunk สังกัด -> แปลงชื่อไฟล์ PDF เป็น "หัวข้อสีชมพูในชีท"
        c_details = []
        doc_pages = {}
        for cid in c_ids_list:
            info = chunk_info_map.get(cid)
            if info:
                src_pdf = info["source"]
                # ดึงชื่อหัวข้อสีชมพูในชีท
                pink_title = pdf_to_pink_title.get(src_pdf, src_pdf)
                page = info["page"]
                if pink_title not in doc_pages:
                    doc_pages[pink_title] = []
                if page not in doc_pages[pink_title]:
                    doc_pages[pink_title].append(page)

        for pink_title, pages in doc_pages.items():
            pages_str = ", ".join(map(str, sorted(pages)))
            c_details.append(f"{pink_title} (หน้า {pages_str})")

        if not c_details and h.get("referenced_docs"):
            try:
                ref_docs = json.loads(h["referenced_docs"])
                for d in ref_docs:
                    pink = pdf_to_pink_title.get(d, d)
                    c_details.append(pink)
            except Exception:
                pass

        c_details_display = "\n".join(c_details) if c_details else "-"

        ws_new.append([
            item_no,
            category,
            q,
            ans,
            exp,
            c_ids_display,
            c_details_display,
            resp_time
        ])

    # 6. จัดแต่งสไตล์ Excel ให้สวยงาม
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # กรมท่าเข้ม
    header_font = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    font_regular = Font(name="Tahoma", size=9, color="1F2937")

    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws_new.row_dimensions[1].height = 30
    for col_idx in range(1, len(headers) + 1):
        cell = ws_new.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for row_idx in range(2, ws_new.max_row + 1):
        ws_new.row_dimensions[row_idx].height = 65
        is_even = (row_idx % 2 == 0)
        row_fill = fill_zebra if is_even else fill_white

        for col_idx in range(1, len(headers) + 1):
            cell = ws_new.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.fill = row_fill
            cell.border = thin_border

            # คอลัมน์ 1: ข้อที่ -> กึ่งกลาง
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            # คอลัมน์ 2: หัวข้อของคำถาม -> กึ่งกลาง/จัดชิดซ้าย
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            # คอลัมน์ 3, 4, 5: คำถาม, คำตอบ, คำตอบคาดหวัง -> ชิดซ้าย
            elif col_idx in (3, 4, 5):
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            # คอลัมน์ 6: Chunk IDs -> กึ่งกลาง
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            # คอลัมน์ 7: เอกสารและหน้า -> ชิดซ้าย
            elif col_idx == 7:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            # คอลัมน์ 8: เวลาในการตอบ -> กึ่งกลาง
            elif col_idx == 8:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00"s"'

    col_widths = {
        1: 8,   # ข้อที่
        2: 30,  # หัวข้อของคำถาม (ในชีท)
        3: 36,  # คำถาม
        4: 48,  # คำตอบที่บอทตอบ
        5: 38,  # คำตอบที่คาดหวัง
        6: 18,  # Chunk ที่ใช้ตอบ (Chunk IDs)
        7: 52,  # หัวข้อเอกสารและหน้าที่ Chunk สังกัด (หัวข้อสีชมพู)
        8: 14   # เวลาในการตอบ
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws_new.column_dimensions[col_letter].width = width

    # บันทึกไฟล์ที่มีชื่อสื่อความหมายชัดเจน
    out_file = CURRENT_DIR / "test_result_gemma_all_80_questions_with_chunks.xlsx"
    pink_file = CURRENT_DIR / "test_result_gemma_with_pink_topics.xlsx"
    wb_new.save(str(pink_file))
    try:
        wb_new.save(str(out_file))
    except PermissionError:
        pass

    try:
        wb_new.save(str(CURRENT_DIR / "test_result_gemma_all_80_questions.xlsx"))
    except PermissionError:
        pass

    print(f"\n🎉 บันทึกไฟล์ผลการทดสอบที่มี Chunk และหัวข้อสีชมพูเรียบร้อยแล้ว:")
    print(f"   -> {pink_file}")
    print(f"   -> {out_file}")


if __name__ == "__main__":
    main()
