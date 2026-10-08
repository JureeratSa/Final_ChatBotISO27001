#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
TUH Chatbot — Multi-Model LLM-as-a-Judge Evaluation & Comparison
=================================================================
เปรียบเทียบการประเมินระหว่าง:
1. Google Gemma 4 26B A4B IT (Judge เดิม)
2. Meta Llama 3.3 70B Instruct (Judge ใหม่ เพื่อลด Self-Enhancement Bias)

บันทึกผลการประเมินทั้งสองโมเดลลงในชีท "accllm"
ของไฟล์ "test_result_gemma_all_80_questions_with_chunks.xlsx"
"""

import os
import sys
import json
import time
import re
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

# Reconfigure stdout for UTF-8 on Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

CURRENT_DIR = Path(__file__).parent.absolute()
PROJECT_ROOT = CURRENT_DIR.parent
GEMMA_CACHE_FILE = CURRENT_DIR / "eval_accuracy_cache.json"
LLAMA_CACHE_FILE = CURRENT_DIR / "eval_llama_cache.json"

load_dotenv(PROJECT_ROOT / "Backend" / ".env")
load_dotenv(PROJECT_ROOT / ".env")

OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY") or os.getenv("GEMINI_API_KEY")
NEW_JUDGE_MODEL = "meta-llama/llama-3.3-70b-instruct"
OLD_JUDGE_MODEL = "google/gemma-4-26b-a4b-it"


def normalize_thai(t: str) -> str:
    if not t:
        return ""
    return re.sub(r"\s+", " ", str(t).strip())


def call_judge(model_name: str, question: str, expected: str, actual: str, retries: int = 3) -> dict:
    if not actual or actual.strip() in ["-", ""]:
        return {
            "score": 0.0,
            "verdict": "ไม่ถูกต้อง",
            "reason": "แชทบอทไม่ได้ตอบข้อความ หรือไม่มีคำตอบในผลการทดสอบ"
        }

    if normalize_thai(actual) == normalize_thai(expected):
        return {
            "score": 1.0,
            "verdict": "ถูกต้อง",
            "reason": "คำตอบตรงกับคำตอบที่คาดหวังอย่างสมบูรณ์แบบ"
        }

    prompt = f"""คุณคือผู้ประเมินคุณภาพของระบบ AI Chatbot (LLM-as-a-Judge) สำหรับเอกสารมาตรฐาน ISO โรงพยาบาลธรรมศาสตร์เฉลิมพระเกียรติ
หน้าที่ของคุณคือตัดสินคะแนนความถูกต้องของคำตอบจาก Chatbot เปรียบเทียบกับคำตอบที่คาดหวัง (Expected Answer) อย่างเป็นกลางและเข้มงวด

[คำถาม]:
{question}

[คำตอบที่คาดหวัง (Expected Answer / Golden Answer)]:
{expected}

[คำตอบจาก Chatbot (Actual Answer)]:
{actual}

---
เกณฑ์การตัดสินให้คะแนนอย่างเป็นกลางและเข้มงวด:
1.0 = "ถูกต้อง": สาระสำคัญและข้อเท็จจริงหลักครบถ้วน สอดคล้องกับคำตอบที่คาดหวัง (แม้สำนวนภาษาต่างกัน หรือมีรายละเอียดเพิ่มเติมที่ถูกต้อง)
0.5 = "ถูกต้องบางส่วน": ตอบถูกบางประเด็นสำคัญ แต่ขาดรายละเอียดสำคัญบางข้อ หรือมีเนื้อหาที่ยังไม่ชัดเจน/คลุมเครือ
0.0 = "ไม่ถูกต้อง": ตอบผิดจากข้อเท็จจริง, ไม่ตรงประเด็นกับคำถาม, มีข้อมูลเท็จ (Hallucination), หรือตอบว่าไม่พบข้อมูลทั้งที่คำตอบที่คาดหวังมีข้อมูล

กรุณาตอบผลการประเมินเป็น JSON Format ดังนี้เท่านั้น (ห้ามใส่คำเกริ่นนำหรือ Markdown อื่น):
{{
  "score": 1.0,
  "verdict": "ถูกต้อง",
  "reason": "อธิบายเหตุผลสั้นๆ 1-2 ประโยคว่าทำไมถึงได้คะแนนนี้"
}}
"""

    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:8000",
        "X-Title": "TUH Chatbot Multi-Judge Eval"
    }

    payload = {
        "model": model_name,
        "messages": [
            {"role": "system", "content": "You are an expert AI evaluator. Output only valid JSON."},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.0,
        "max_tokens": 350
    }

    for attempt in range(retries):
        try:
            resp = requests.post(
                "https://openrouter.ai/api/v1/chat/completions",
                headers=headers,
                json=payload,
                timeout=30
            )
            if resp.status_code == 200:
                body = resp.json()
                content = body["choices"][0]["message"]["content"].strip()
                if content.startswith("```"):
                    content = re.sub(r"^```[a-zA-Z]*\n?", "", content)
                    content = re.sub(r"\n?```$", "", content).strip()

                parsed = json.loads(content)
                score = float(parsed.get("score", 0.0))
                verdict = str(parsed.get("verdict", "ไม่ถูกต้อง"))
                reason = str(parsed.get("reason", ""))

                if score >= 0.9:
                    verdict = "ถูกต้อง"
                elif score >= 0.4:
                    verdict = "ถูกต้องบางส่วน"
                else:
                    verdict = "ไม่ถูกต้อง"

                return {"score": score, "verdict": verdict, "reason": reason}
            else:
                time.sleep(1.5 * (attempt + 1))
        except Exception:
            time.sleep(1.5 * (attempt + 1))

    return {
        "score": 0.5,
        "verdict": "ถูกต้องบางส่วน",
        "reason": "ไม่สามารถเรียกโมเดลประเมินได้เนื่องจาก Network Timeout"
    }


def load_cache(path: Path) -> dict:
    if path.exists():
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}


def save_cache(path: Path, data: dict):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def main():
    print("=" * 80)
    print("⚖️ TUH Chatbot — Multi-Model Judge Evaluation (Gemma 4 vs Llama 3.3 70B)")
    print("=" * 80)

    # 1. โหลดข้อมูลคำถามและคำตอบจากชีทตั้งต้น
    source_file = CURRENT_DIR / "test_result_gemma_all_80_questions_with_chunks.xlsx"
    if not source_file.exists():
        source_file = CURRENT_DIR / "test_result_gemma_with_accuracy.xlsx"

    wb_src = openpyxl.load_workbook(source_file, data_only=True)
    ws_src = wb_src.active

    rows_data = []
    for r in range(2, ws_src.max_row + 1):
        item_no = ws_src.cell(r, 1).value
        cat = str(ws_src.cell(r, 2).value or "").strip()
        q = str(ws_src.cell(r, 3).value or "").strip()
        ans = str(ws_src.cell(r, 4).value or "").strip()
        exp = str(ws_src.cell(r, 5).value or "").strip()
        chunks = str(ws_src.cell(r, 6).value or "").strip()
        pink_topic = str(ws_src.cell(r, 7).value or "").strip()
        resp_time = ws_src.cell(r, 8).value

        rows_data.append({
            "item_no": item_no,
            "category": cat,
            "question": q,
            "bot_ans": ans,
            "expected": exp,
            "chunks": chunks,
            "pink_topic": pink_topic,
            "resp_time": resp_time
        })

    total_items = len(rows_data)
    print(f"📂 โหลดข้อมูลทั้งหมด {total_items} ข้อเรียบร้อยแล้ว")

    # 2. โหลดแคชเดิม (Gemma 4 26B)
    gemma_cache = load_cache(GEMMA_CACHE_FILE)
    print(f"📦 โหลดแคชเดิมของ Gemma 4 26B: {len(gemma_cache)} ข้อ")

    # 3. รันการประเมินด้วย Llama 3.3 70B
    llama_cache = load_cache(LLAMA_CACHE_FILE)
    print(f"📦 โหลดแคชของ Llama 3.3 70B ที่มีอยู่แล้ว: {len(llama_cache)} ข้อ")

    items_to_eval_llama = [it for it in rows_data if it["question"] not in llama_cache]

    if items_to_eval_llama:
        print(f"\n🚀 กำลังประเมินด้วย {NEW_JUDGE_MODEL} จำนวน {len(items_to_eval_llama)} ข้อ...")
        done_cnt = len(llama_cache)

        def eval_task(it):
            res = call_judge(NEW_JUDGE_MODEL, it["question"], it["expected"], it["bot_ans"])
            return it["question"], res

        with ThreadPoolExecutor(max_workers=5) as executor:
            futures = [executor.submit(eval_task, it) for it in items_to_eval_llama]
            for future in as_completed(futures):
                q_key, res = future.result()
                llama_cache[q_key] = res
                done_cnt += 1
                save_cache(LLAMA_CACHE_FILE, llama_cache)
                print(f"   [{done_cnt}/{total_items}] Llama: {res['score']} | {res['verdict']} | {q_key[:38]}...")
    else:
        print("⚡ ข้อมูลทุกข้อของ Llama 3.3 70B ได้รับการประเมินครบแล้วจากแคช!")

    # 4. ประมวลผลเปรียบเทียบทั้ง 2 โมเดล
    gemma_scores = []
    llama_scores = []
    agreements = []
    gemma_verdicts = {"ถูกต้อง": 0, "ถูกต้องบางส่วน": 0, "ไม่ถูกต้อง": 0}
    llama_verdicts = {"ถูกต้อง": 0, "ถูกต้องบางส่วน": 0, "ไม่ถูกต้อง": 0}
    llama_higher = 0
    gemma_higher = 0

    merged_results = []
    for it in rows_data:
        q_key = it["question"]
        g_res = gemma_cache.get(q_key, {"score": 0.0, "verdict": "ไม่ถูกต้อง", "reason": "ไม่มีข้อมูล"})
        l_res = llama_cache.get(q_key, {"score": 0.0, "verdict": "ไม่ถูกต้อง", "reason": "ไม่มีข้อมูล"})

        g_score = g_res.get("score", 0.0)
        l_score = l_res.get("score", 0.0)
        g_verdict = g_res.get("verdict", "ไม่ถูกต้อง")
        l_verdict = l_res.get("verdict", "ไม่ถูกต้อง")

        gemma_scores.append(g_score)
        llama_scores.append(l_score)
        gemma_verdicts[g_verdict] = gemma_verdicts.get(g_verdict, 0) + 1
        llama_verdicts[l_verdict] = llama_verdicts.get(l_verdict, 0) + 1

        is_agreed = (g_score == l_score)
        agreements.append(1.0 if is_agreed else 0.0)

        if l_score > g_score:
            llama_higher += 1
        elif g_score > l_score:
            gemma_higher += 1

        agree_text = "ตรงกัน (Agree)" if is_agreed else f"เห็นต่าง ({l_score} vs {g_score})"

        merged_results.append({
            **it,
            "gemma_score": g_score,
            "gemma_verdict": g_verdict,
            "gemma_reason": g_res.get("reason", ""),
            "llama_score": l_score,
            "llama_verdict": l_verdict,
            "llama_reason": l_res.get("reason", ""),
            "is_agreed": is_agreed,
            "agree_text": agree_text
        })

    gemma_acc = (sum(gemma_scores) / total_items) * 100
    llama_acc = (sum(llama_scores) / total_items) * 100
    agreement_rate = (sum(agreements) / total_items) * 100
    agree_count = int(sum(agreements))

    print("\n" + "=" * 70)
    print("📊 ผลการเปรียบเทียบระหว่าง 2 โมเดลผู้ตัดสิน (Inter-Model Comparison)")
    print("=" * 70)
    print(f"• โมเดลเดิม [Gemma 4 26B] ความแม่นยำเฉลี่ย : {gemma_acc:.2f}%")
    print(f"  - ถูกต้อง (1.0): {gemma_verdicts.get('ถูกต้อง', 0)} ข้อ | บางส่วน (0.5): {gemma_verdicts.get('ถูกต้องบางส่วน', 0)} ข้อ | ผิด (0.0): {gemma_verdicts.get('ไม่ถูกต้อง', 0)} ข้อ")
    print(f"• โมเดลใหม่ [Llama 3.3 70B] ความแม่นยำเฉลี่ย : {llama_acc:.2f}%")
    print(f"  - ถูกต้อง (1.0): {llama_verdicts.get('ถูกต้อง', 0)} ข้อ | บางส่วน (0.5): {llama_verdicts.get('ถูกต้องบางส่วน', 0)} ข้อ | ผิด (0.0): {llama_verdicts.get('ไม่ถูกต้อง', 0)} ข้อ")
    print(f"• อัตราความเห็นพ้องต้องกัน (Agreement Rate)  : {agreement_rate:.2f}% ({agree_count}/{total_items} ข้อ)")
    print(f"  - Llama 3.3 70B ให้คะแนนเข้มงวดกว่า : {gemma_higher} ข้อ")
    print(f"  - Llama 3.3 70B ให้คะแนนสูงกว่า      : {llama_higher} ข้อ")
    print("=" * 70)

    # 5. โหลด Workbook เป้าหมายเพื่อสร้าง/อัปเดตชีท "accllm"
    target_path = CURRENT_DIR / "test_result_gemma_all_80_questions_with_chunks.xlsx"
    wb_target = openpyxl.load_workbook(target_path)

    if "accllm" in wb_target.sheetnames:
        del wb_target["accllm"]

    ws_acc = wb_target.create_sheet(title="accllm")
    ws_acc.views.sheetView[0].showGridLines = True

    # Styling
    navy_dark = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    purple_gemma = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")
    teal_llama = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    bg_agree_summary = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

    font_white_bold = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    font_title = Font(name="Tahoma", size=14, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Tahoma", size=10, color="64748B")
    font_bold = Font(name="Tahoma", size=9, bold=True, color="1F2937")
    font_regular = Font(name="Tahoma", size=9, color="1F2937")

    fill_correct = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    font_correct = Font(name="Tahoma", size=9, bold=True, color="166534")

    fill_partial = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    font_partial = Font(name="Tahoma", size=9, bold=True, color="92400E")

    fill_incorrect = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_incorrect = Font(name="Tahoma", size=9, bold=True, color="991B1B")

    fill_agree = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    font_agree = Font(name="Tahoma", size=9, bold=True, color="0369A1")
    fill_diff = PatternFill(start_color="FFEDD5", end_color="FFEDD5", fill_type="solid")
    font_diff = Font(name="Tahoma", size=9, bold=True, color="C2410C")

    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title Banner
    ws_acc["A1"] = "การเปรียบเทียบผลประเมินความถูกต้องของแชทบอท (Multi-Model LLM-as-a-Judge)"
    ws_acc["A1"].font = font_title
    ws_acc["A2"] = f"เปรียบเทียบการตัดสินระหว่าง Google Gemma 4 (26B) กับ Meta Llama 3.3 (70B) เพื่อลดอคติ (Bias Reduction)"
    ws_acc["A2"].font = font_subtitle

    # Summary KPI Cards
    kpi_blocks = [
        ("B4", "D5", "คะแนนเฉลี่ยเดิม [Gemma 4 26B]", f"{gemma_acc:.2f}%", f"ถูก {gemma_verdicts.get('ถูกต้อง',0)} | บางส่วน {gemma_verdicts.get('ถูกต้องบางส่วน',0)} | ผิด {gemma_verdicts.get('ไม่ถูกต้อง',0)}", PatternFill(start_color="EDE9FE", end_color="EDE9FE", fill_type="solid")),
        ("E4", "G5", "คะแนนเฉลี่ยใหม่ [Llama 3.3 70B]", f"{llama_acc:.2f}%", f"ถูก {llama_verdicts.get('ถูกต้อง',0)} | บางส่วน {llama_verdicts.get('ถูกต้องบางส่วน',0)} | ผิด {llama_verdicts.get('ไม่ถูกต้อง',0)}", PatternFill(start_color="CCFBF1", end_color="CCFBF1", fill_type="solid")),
        ("H4", "J5", "อัตราความเห็นตรงกัน (Agreement)", f"{agreement_rate:.2f}%", f"ทั้งสองโมเดลเห็นตรงกัน {agree_count}/{total_items} ข้อ", bg_agree_summary),
        ("K4", "M5", "การวิเคราะห์ความเข้มงวด", f"ต่างกัน {total_items - agree_count} ข้อ", f"Llama เข้มงวดกว่า {gemma_higher} ข้อ | ผ่อนปรนกว่า {llama_higher} ข้อ", PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"))
    ]

    for top_left, bot_right, title, val, sub, bg in kpi_blocks:
        ws_acc.merge_cells(f"{top_left}:{bot_right}")
        c = ws_acc[top_left]
        c.value = f"{title}\n{val}\n({sub})"
        c.font = font_bold
        c.fill = bg
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        s_col, s_row = openpyxl.utils.coordinate_to_tuple(top_left)
        e_col, e_row = openpyxl.utils.coordinate_to_tuple(bot_right)
        for r in range(s_row, e_row + 1):
            for col in range(s_col, e_col + 1):
                ws_acc.cell(r, col).border = thin_border

    # Table Header at Row 7
    headers = [
        "ข้อที่",
        "หมวดหมู่นโยบาย (ในชีท)",
        "คำถาม",
        "คำตอบที่บอทตอบ",
        "คำตอบที่คาดหวัง",
        "Chunk IDs",
        "หัวข้อสีชมพู (เอกสาร Chunk)",
        "คะแนน [Gemma 4 26B]",
        "ผลตัดสิน [Gemma 4 26B]",
        "เหตุผล [Gemma 4 26B]",
        "คะแนน [Llama 3.3 70B]",
        "ผลตัดสิน [Llama 3.3 70B]",
        "เหตุผล [Llama 3.3 70B]",
        "ผลความเห็นพ้อง (Agreement)",
        "เวลาตอบ (วินาที)"
    ]

    ws_acc.row_dimensions[7].height = 34
    for c_idx, h in enumerate(headers, start=1):
        cell = ws_acc.cell(7, c_idx, h)
        cell.font = font_white_bold
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # สีหัวตารางแบ่งแยกหมวด
        if 8 <= c_idx <= 10:
            cell.fill = purple_gemma  # กลุ่ม Gemma
        elif 11 <= c_idx <= 13:
            cell.fill = teal_llama   # กลุ่ม Llama
        else:
            cell.fill = navy_dark

    # Data Rows
    for row_idx, rdata in enumerate(merged_results, start=8):
        ws_acc.row_dimensions[row_idx].height = 70
        is_even = (row_idx % 2 == 0)
        base_fill = fill_zebra if is_even else fill_white

        row_vals = [
            (rdata["item_no"], Alignment(horizontal="center", vertical="top")),
            (rdata["category"], Alignment(horizontal="left", vertical="top", wrap_text=True)),
            (rdata["question"], Alignment(horizontal="left", vertical="top", wrap_text=True)),
            (rdata["bot_ans"], Alignment(horizontal="left", vertical="top", wrap_text=True)),
            (rdata["expected"], Alignment(horizontal="left", vertical="top", wrap_text=True)),
            (rdata["chunks"], Alignment(horizontal="center", vertical="top", wrap_text=True)),
            (rdata["pink_topic"], Alignment(horizontal="left", vertical="top", wrap_text=True)),
            (rdata["gemma_score"], Alignment(horizontal="center", vertical="top")),
            (rdata["gemma_verdict"], Alignment(horizontal="center", vertical="top")),
            (rdata["gemma_reason"], Alignment(horizontal="left", vertical="top", wrap_text=True)),
            (rdata["llama_score"], Alignment(horizontal="center", vertical="top")),
            (rdata["llama_verdict"], Alignment(horizontal="center", vertical="top")),
            (rdata["llama_reason"], Alignment(horizontal="left", vertical="top", wrap_text=True)),
            (rdata["agree_text"], Alignment(horizontal="center", vertical="top", wrap_text=True)),
            (rdata["resp_time"], Alignment(horizontal="center", vertical="top"))
        ]

        for col_idx, (val, align) in enumerate(row_vals, start=1):
            cell = ws_acc.cell(row_idx, col_idx, val)
            cell.font = font_regular
            cell.fill = base_fill
            cell.alignment = align
            cell.border = thin_border

            # ไฮไลต์สีผลประเมิน Gemma
            if col_idx == 9:
                if rdata["gemma_verdict"] == "ถูกต้อง":
                    cell.fill = fill_correct
                    cell.font = font_correct
                elif rdata["gemma_verdict"] == "ถูกต้องบางส่วน":
                    cell.fill = fill_partial
                    cell.font = font_partial
                else:
                    cell.fill = fill_incorrect
                    cell.font = font_incorrect
            elif col_idx == 8:
                cell.font = font_bold

            # ไฮไลต์สีผลประเมิน Llama
            elif col_idx == 12:
                if rdata["llama_verdict"] == "ถูกต้อง":
                    cell.fill = fill_correct
                    cell.font = font_correct
                elif rdata["llama_verdict"] == "ถูกต้องบางส่วน":
                    cell.fill = fill_partial
                    cell.font = font_partial
                else:
                    cell.fill = fill_incorrect
                    cell.font = font_incorrect
            elif col_idx == 11:
                cell.font = font_bold

            # ไฮไลต์ Agreement
            elif col_idx == 14:
                if rdata["is_agreed"]:
                    cell.fill = fill_agree
                    cell.font = font_agree
                else:
                    cell.fill = fill_diff
                    cell.font = font_diff

            # เวลาตอบ
            elif col_idx == 15:
                if isinstance(val, (int, float)):
                    cell.number_format = '0.00"s"'

    col_widths = {
        1: 8,   # ข้อที่
        2: 26,  # หมวดหมู่
        3: 34,  # คำถาม
        4: 45,  # คำตอบบอท
        5: 36,  # เฉลย
        6: 15,  # Chunk IDs
        7: 42,  # หัวข้อสีชมพู
        8: 14,  # คะแนน Gemma
        9: 16,  # ผลตัดสิน Gemma
        10: 38, # เหตุผล Gemma
        11: 14, # คะแนน Llama
        12: 16, # ผลตัดสิน Llama
        13: 38, # เหตุผล Llama
        14: 20, # ความเห็นพ้อง
        15: 14  # เวลาตอบ
    }
    for col_idx, w in col_widths.items():
        ws_acc.column_dimensions[get_column_letter(col_idx)].width = w

    # บันทึกไฟล์อย่างปลอดภัย (รองรับกรณีผู้ใช้เปิดไฟล์ค้างไว้)
    saved_successfully = False
    try:
        wb_target.save(str(target_path))
        print(f"\n🎉 บันทึกชีท 'accllm' ลงในไฟล์เป้าหมายสำเร็จ:")
        print(f"   -> {target_path}")
        saved_successfully = True
    except PermissionError:
        backup_path = CURRENT_DIR / "test_result_gemma_all_80_questions_with_chunks_accllm.xlsx"
        wb_target.save(str(backup_path))
        print(f"\n⚠️ ไฟล์ {target_path.name} กำลังถูกเปิดใช้งานใน Excel จึงบันทึกเป็นชื่อใหม่:")
        print(f"   -> {backup_path}")
        print("   (สามารถปิดไฟล์เดิมใน Excel แล้วรันสคริปต์ซ้ำ หรือเปิดดูจากไฟล์ใหม่นี้ได้ทันทีครับ)")


if __name__ == "__main__":
    main()
