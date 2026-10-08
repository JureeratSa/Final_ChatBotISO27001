#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
TUH Chatbot AI — Automated Test Runner & Excel Report Generator
โฟลเดอร์: testchatbotSPO/
===============================================================
สคริปต์ทดสอบระบบแชทบอท TUH:
1. ดึงชุดคำถามจาก Google Sheets URL (หรือ fallback ไฟล์ในเครื่องอัตโนมัติหากติดสิทธิ์ 401)
2. สุ่มลำดับคำถาม (random.shuffle) ตามที่ผู้ใช้งานกำหนด
3. ตรวจสอบและบังคับใช้โมเดล google/gemma-4-26b-a4b-it
4. ส่งคำถามเข้า Backend API (http://localhost:8000/api/chat) เก็บคำตอบจริงและเวลาตอบสนอง
5. สร้างไฟล์ผลการทดสอบเป็น Excel (.xlsx) สวยงามพร้อมคอลัมน์:
   - ข้อที่
   - คำถาม
   - คำตอบที่บอทตอบ
   - คำตอบที่คาดหวัง
   - เวลาในการตอบ

การใช้งาน:
    py run_test.py
    py run_test.py --limit 10
    py run_test.py --file questions_dataset.csv
    py run_test.py --sheet "https://docs.google.com/spreadsheets/d/.../edit"
"""

import os
import sys
import time
import io
import json
import random
import re
import argparse
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

import requests
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# บังคับใช้ UTF-8 stdout บน Windows Console
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

# กำหนดค่าเริ่มต้น
DEFAULT_SHEET_URL = "https://docs.google.com/spreadsheets/d/1M3KaHBdxmUsE6zg716IW-QmpWe32hNxUxUF6T62SGZQ/edit?usp=sharing"
DEFAULT_API_URL = "http://localhost:8000/api/chat"
TARGET_MODEL = "google/gemma-4-26b-a4b-it"
DEFAULT_DELAY = 2.0  # วินาที เพื่อป้องกัน Rate Limit (30 req / 60s)
DEFAULT_TIMEOUT = 90  # วินาทีต่อคำถาม
CURRENT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = CURRENT_DIR.parent


def extract_google_sheet_id(url: str) -> Optional[str]:
    """ดึง Spreadsheet ID จาก URL ของ Google Sheets"""
    match = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url)
    return match.group(1) if match else None


def fetch_from_google_sheet(sheet_url: str) -> Tuple[Optional[pd.DataFrame], Optional[str]]:
    """
    พยายามดาวน์โหลดข้อมูลจาก Google Sheets ผ่าน Export URL
    คืนค่า (DataFrame, ErrorMessage)
    """
    sheet_id = extract_google_sheet_id(sheet_url)
    if not sheet_id:
        return None, "รูปแบบ URL ของ Google Sheets ไม่ถูกต้อง"

    export_urls = [
        (f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx", "xlsx"),
        (f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv", "csv"),
        (f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv", "csv"),
    ]

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    }

    last_error = None
    for url, fmt in export_urls:
        try:
            resp = requests.get(url, headers=headers, timeout=15, allow_redirects=True)
            if resp.status_code == 200 and len(resp.content) > 100:
                content_type = resp.headers.get("Content-Type", "").lower()
                if "html" in content_type and ("accounts.google.com" in resp.text or "Sign in" in resp.text):
                    continue

                if fmt == "xlsx":
                    df = pd.read_excel(io.BytesIO(resp.content))
                    return df, None
                elif fmt == "csv":
                    try:
                        text = resp.content.decode("utf-8-sig")
                    except UnicodeDecodeError:
                        text = resp.content.decode("cp874", errors="ignore")
                    df = pd.read_csv(io.StringIO(text))
                    return df, None

            elif resp.status_code == 401:
                last_error = (
                    "Google Sheets ตอบกลับ HTTP 401 Unauthorized (ลิงก์ถูกจำกัดสิทธิ์เข้าถึงเฉพาะผู้ที่ล็อกอิน)\n"
                    "  💡 วิธีเปิดสิทธิ์: ใน Google Sheet กด 'แชร์ (Share)' -> 'การเข้าถึงทั่วไป' -> เปลี่ยนเป็น 'ทุกคนที่มีลิงก์ (Anyone with the link)'"
                )
            else:
                last_error = f"HTTP Status {resp.status_code}"
        except Exception as e:
            last_error = str(e)

    return None, last_error


def load_local_file(file_path: Path) -> Tuple[Optional[pd.DataFrame], Optional[str]]:
    """โหลดข้อมูลชุดคำถามจากไฟล์ในเครื่อง (.xlsx, .xls, .csv)"""
    if not file_path.exists():
        return None, f"ไม่พบไฟล์: {file_path}"

    try:
        suffix = file_path.suffix.lower()
        if suffix in [".xlsx", ".xls"]:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet_name = "คำถาม" if "คำถาม" in wb.sheetnames else wb.sheetnames[0]
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            return df, None
        elif suffix == ".csv":
            try:
                df = pd.read_csv(file_path, encoding="utf-8-sig")
            except Exception:
                df = pd.read_csv(file_path, encoding="cp874")
            return df, None
        else:
            return None, f"ไม่รองรับนามสกุลไฟล์: {suffix} (รองรับเฉพาะ .xlsx, .xls, .csv)"
    except Exception as e:
        return None, f"เกิดข้อผิดพลาดในการอ่านไฟล์ {file_path}: {e}"


def normalize_questions(df: pd.DataFrame) -> List[Dict[str, Any]]:
    """
    ตรวจจับชื่อคอลัมน์อัตโนมัติ:
    - คำถาม
    - คำตอบที่คาดหวัง
    - ข้อที่ / ลำดับ
    """
    cols = list(df.columns)
    q_col = None
    a_col = None
    no_col = None

    cat_col = None
    q_patterns = ["คำถาม", "question", "query", "ข้อความ", "prompt"]
    a_patterns = ["คำตอบที่คาดหวัง", "คำตอบที่คาดว่าจะได้รับ", "expected_answer", "expected", "คำตอบ", "answer"]
    no_patterns = ["ข้อที่", "ลำดับ", "id", "no", "number", "ข้อ"]
    cat_patterns = ["หัวข้อ", "หมวดหมู่", "category", "หัวข้อคำถาม", "หมวด"]

    for col in cols:
        c_clean = str(col).strip().lower()
        if not q_col and any(p in c_clean for p in q_patterns):
            q_col = col
        elif not a_col and any(p in c_clean for p in a_patterns):
            a_col = col
        elif not no_col and any(p in c_clean for p in no_patterns):
            no_col = col
        elif not cat_col and any(p in c_clean for p in cat_patterns):
            cat_col = col

    # Fallback กรณีหัวตารางไม่ตรง
    if not q_col and len(cols) >= 2:
        q_col = cols[1] if len(cols) > 1 else cols[0]
    if not a_col and len(cols) >= 3:
        a_col = cols[2]

    records = []
    for idx, row in df.iterrows():
        question = str(row[q_col]).strip() if q_col and pd.notna(row[q_col]) else ""
        if not question or question.lower() == "nan":
            continue

        expected = str(row[a_col]).strip() if a_col and pd.notna(row[a_col]) and str(row[a_col]).lower() != "nan" else "-"
        item_no = str(row[no_col]).strip() if no_col and pd.notna(row[no_col]) and str(row[no_col]).lower() != "nan" else str(idx + 1)
        category = str(row[cat_col]).strip() if cat_col and pd.notna(row[cat_col]) and str(row[cat_col]).lower() != "nan" else "-"

        records.append({
            "orig_no": item_no,
            "category": category,
            "question": question,
            "expected_answer": expected
        })

    return records


def verify_and_set_model_in_db():
    """ตรวจสอบและอัปเดตโมเดลในฐานข้อมูลให้เป็น google/gemma-4-26b-a4b-it"""
    try:
        import asyncio
        if sys.platform == "win32":
            asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

        backend_dir = PROJECT_ROOT / "Backend"
        if str(backend_dir) not in sys.path:
            sys.path.insert(0, str(backend_dir.resolve()))

        from app.core.database import AsyncSessionLocal
        from app.models.models import SystemSettings
        from sqlalchemy import select

        async def _update():
            async with AsyncSessionLocal() as session:
                res = await session.execute(select(SystemSettings).where(SystemSettings.id == "config"))
                cfg = res.scalar_one_or_none()
                if cfg:
                    if cfg.model_name != TARGET_MODEL:
                        print(f"🔄 กำลังปรับโมเดลในฐานข้อมูลจาก '{cfg.model_name}' -> '{TARGET_MODEL}'...")
                        cfg.model_name = TARGET_MODEL
                        await session.commit()
                        print(f"✅ ปรับโมเดลในฐานข้อมูลเป็น '{TARGET_MODEL}' เรียบร้อยแล้ว")
                    else:
                        print(f"✅ โมเดลในฐานข้อมูลตั้งเป็น '{TARGET_MODEL}' อยู่แล้ว")

        asyncio.run(_update())
    except Exception as e:
        print(f"⚠️ ไม่สามารถอัปเดตโมเดลใน DB อัตโนมัติ (อาจรันต่อได้): {e}")


def query_chatbot(api_url: str, question: str, timeout: int = DEFAULT_TIMEOUT) -> Dict[str, Any]:
    """ส่งคำถามเข้า Chatbot API และวัดเวลาการตอบสนอง พร้อมรองรับ retry เมื่อชน rate limit"""
    max_retries = 3
    retry_delay = 10

    for attempt in range(max_retries):
        start_time = time.perf_counter()
        try:
            resp = requests.post(
                api_url,
                json={"query": question},
                headers={"Content-Type": "application/json; charset=utf-8"},
                timeout=timeout
            )
            elapsed = time.perf_counter() - start_time

            if resp.status_code == 200:
                data = resp.json()
                bot_answer = data.get("answer", "")
                bot_answer = bot_answer.replace("__API_URL__", "http://localhost:8000")
                model_used = data.get("model", "")
                return {
                    "answer": bot_answer.strip(),
                    "elapsed": round(elapsed, 2),
                    "model": model_used,
                    "used_rag": data.get("used_rag", False),
                    "error": None
                }
            elif resp.status_code == 429:
                print(f"\n   ⚠️ ติด Rate Limit (HTTP 429) พักรอ {retry_delay} วินาทีแล้วลองใหม่ (ครั้งที่ {attempt+1}/{max_retries})...")
                time.sleep(retry_delay)
                retry_delay += 5
                continue
            else:
                return {
                    "answer": f"[HTTP {resp.status_code}] {resp.text[:200]}",
                    "elapsed": round(elapsed, 2),
                    "model": "error",
                    "error": f"HTTP {resp.status_code}"
                }
        except requests.exceptions.ConnectionError:
            return {
                "answer": "[Connection Error] ไม่สามารถเชื่อมต่อกับ Backend ได้ (กรุณารัน python run_backend.py ก่อน)",
                "elapsed": 0.0,
                "model": "error",
                "error": "ConnectionError"
            }
        except requests.exceptions.Timeout:
            return {
                "answer": f"[Timeout] บอทไม่ตอบสนองภายใน {timeout} วินาที",
                "elapsed": round(time.perf_counter() - start_time, 2),
                "model": "error",
                "error": "Timeout"
            }
        except Exception as e:
            return {
                "answer": f"[Error] {str(e)}",
                "elapsed": round(time.perf_counter() - start_time, 2),
                "model": "error",
                "error": str(e)
            }

    return {
        "answer": "[Rate Limit] ส่งคำขอถี่เกินไปแม้ลองใหม่แล้ว",
        "elapsed": 0.0,
        "model": "error",
        "error": "HTTP 429"
    }


def export_to_styled_excel(results: List[Dict[str, Any]], output_path: Path):
    """
    บันทึกผลการทดสอบลงไฟล์ Excel (.xlsx) และจัดสไตล์อย่างสวยงาม
    คอลัมน์:
    1. ข้อที่
    2. คำถาม
    3. คำตอบที่บอทตอบ
    4. คำตอบที่คาดหวัง
    5. เวลาในการตอบ
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ผลการทดสอบแชทบอท"

    headers = [
        "ข้อที่",
        "คำถาม",
        "คำตอบที่บอทตอบ",
        "คำตอบที่คาดหวัง",
        "เวลาในการตอบ"
    ]
    ws.append(headers)

    for item in results:
        ws.append([
            item["item_no"],
            item["question"],
            item["bot_answer"],
            item["expected_answer"],
            item["response_time"]
        ])

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")  # Navy Blue
    header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Tahoma", size=10, color="1F2937")

    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for row_idx in range(2, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 50
        is_even = (row_idx % 2 == 0)
        row_fill = fill_zebra if is_even else fill_white

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.fill = row_fill
            cell.border = thin_border

            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            elif col_idx in (2, 3, 4):
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="center", vertical="top")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.00"s"'

    col_widths = {
        1: 10,  # ข้อที่
        2: 38,  # คำถาม
        3: 54,  # คำตอบที่บอทตอบ
        4: 44,  # คำตอบที่คาดหวัง
        5: 16   # เวลาในการตอบ
    }
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    print(f"\n🎉 บันทึกผลการทดสอบลง Excel เรียบร้อยแล้วที่: {output_path}")


def main():
    parser = argparse.ArgumentParser(description="TUH Chatbot Test Runner (Gemma 4 26B A4B IT)")
    parser.add_argument("--sheet", type=str, default=DEFAULT_SHEET_URL, help="URL Google Sheets")
    parser.add_argument("--file", type=str, default=None, help="ไฟล์คำถามในเครื่อง (.xlsx, .csv)")
    parser.add_argument("--api", type=str, default=DEFAULT_API_URL, help="URL Chatbot API (default: http://localhost:8000/api/chat)")
    parser.add_argument("--limit", type=int, default=None, help="จำกัดจำนวนข้อที่จะเทส (เช่น --limit 10)")
    parser.add_argument("--delay", type=float, default=DEFAULT_DELAY, help="หน่วงเวลาระหว่างข้อ (วินาที)")
    parser.add_argument("--no-shuffle", action="store_true", help="ปิดการสุ่มลำดับคำถาม")
    parser.add_argument("--output", type=str, default=None, help="ชื่อไฟล์ผลลัพธ์ Excel (.xlsx)")
    args = parser.parse_args()

    print("=" * 75)
    print("🤖 TUH Chatbot AI — Automated Test Runner")
    print(f"🎯 โมเดลเป้าหมาย: {TARGET_MODEL}")
    print("=" * 75)

    # ตรวจสอบและตั้งค่าโมเดลใน DB
    verify_and_set_model_in_db()

    # 1. โหลดข้อมูลคำถาม
    df = None
    load_err = None

    if args.file:
        custom_p = Path(args.file)
        if not custom_p.is_absolute():
            custom_p = CURRENT_DIR / custom_p
        print(f"📂 กำลังโหลดคำถามจากไฟล์ระบุ: {custom_p}")
        df, load_err = load_local_file(custom_p)
    else:
        print(f"🌐 กำลังดึงข้อมูลชุดคำถามจาก Google Sheets...")
        df, load_err = fetch_from_google_sheet(args.sheet)

        if df is None:
            print(f"\n⚠️ {load_err}")
            # Fallback ไปยังไฟล์ในโฟลเดอร์ testchatbotSPO
            fallback_candidates = [
                CURRENT_DIR / "questions_dataset.csv",
                CURRENT_DIR / "questions_dataset.xlsx",
                PROJECT_ROOT / "TestPFM" / "eval" / "golden_qa.csv",
            ]
            for fb in fallback_candidates:
                if fb.exists():
                    print(f"🔄 ตรวจพบไฟล์ชุดคำถามสำรอง: {fb.name} -> ดึงคำถามจากไฟล์นี้แทน...")
                    df, load_err = load_local_file(fb)
                    if df is not None:
                        break

    if df is None:
        print(f"❌ ไม่สามารถโหลดชุดคำถามได้: {load_err}")
        sys.exit(1)

    test_cases = normalize_questions(df)
    if not test_cases:
        print("❌ ไม่พบคำถามในชุดข้อมูล")
        sys.exit(1)

    print(f"✅ โหลดคำถามสำเร็จทั้งหมด {len(test_cases)} ข้อ")

    # 2. สุ่มลำดับคำถาม (ตามที่ผู้ใช้สั่ง 'ไม่ต้องเรียงเอาเเบบสุ่มๆ')
    if not args.no_shuffle:
        print("🔀 สุ่มลำดับคำถามเรียบร้อยแล้ว (Random Shuffle)")
        random.shuffle(test_cases)
    else:
        print("➡️ ส่งคำถามตามลำดับเดิม (No Shuffle)")

    # ตัดจำนวนตาม limit ถ้ามี
    if args.limit and args.limit > 0:
        test_cases = test_cases[:args.limit]
        print(f"✂️ จำกัดการทดสอบไว้ที่ {len(test_cases)} ข้อ")

    # 3. ตรวจสอบการเชื่อมต่อ API
    print(f"🔌 ตรวจสอบการเชื่อมต่อกับ Backend ที่: {args.api}")
    health = query_chatbot(args.api, "สวัสดีครับ", timeout=15)
    if health.get("error") == "ConnectionError":
        print("\n❌ ไม่สามารถเชื่อมต่อกับ Backend ได้!")
        print("👉 กรุณาเปิด Terminal อีกหน้าต่างแล้วสั่งรัน:")
        print("     python run_backend.py")
        sys.exit(1)
    else:
        print(f"✅ เชื่อมต่อ Backend สำเร็จ (โมเดลตอบกลับ: {health.get('model')})\n")

    # 4. เริ่มส่งคำถามจริง
    results = []
    times = []
    total = len(test_cases)

    print("-" * 75)
    print(f"🚀 เริ่มต้นการทดสอบจริง {total} ข้อ (หน่วงเวลา {args.delay}s เพื่อเลี่ยง Rate Limit)...")
    print("-" * 75)

    for i, item in enumerate(test_cases, 1):
        q = item["question"]
        expected = item["expected_answer"]

        q_disp = (q[:50] + "...") if len(q) > 50 else q
        sys.stdout.write(f"[{i:02d}/{total:02d}] คำถาม: {q_disp:<54}")
        sys.stdout.flush()

        res = query_chatbot(args.api, q)
        ans = res["answer"]
        elapsed = res["elapsed"]
        times.append(elapsed)

        results.append({
            "item_no": i,
            "question": q,
            "bot_answer": ans,
            "expected_answer": expected,
            "response_time": elapsed
        })

        ans_disp = (ans.replace('\n', ' ')[:35] + "...") if len(ans) > 35 else ans.replace('\n', ' ')
        print(f" ⏱️ {elapsed:>5.2f}s | {ans_disp}")

        if i < total and args.delay > 0:
            time.sleep(args.delay)

    # 5. สรุปผล
    avg_t = sum(times) / len(times) if times else 0
    max_t = max(times) if times else 0
    min_t = min(times) if times else 0

    print("\n" + "=" * 75)
    print("📊 สรุปผลการทดสอบ")
    print(f"  • จำนวนคำถามทั้งหมด: {len(results)} ข้อ")
    print(f"  • เวลาตอบเฉลี่ย:      {avg_t:.2f} วินาที")
    print(f"  • เร็วที่สุด / ช้าที่สุด:  {min_t:.2f}s / {max_t:.2f}s")
    print("=" * 75)

    # 6. บันทึกผลลัพธ์ลง Excel
    if not args.output:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_file = CURRENT_DIR / f"test_result_gemma_{ts}.xlsx"
    else:
        out_file = Path(args.output)
        if not out_file.is_absolute():
            out_file = CURRENT_DIR / out_file

    export_to_styled_excel(results, out_file)


if __name__ == "__main__":
    main()
